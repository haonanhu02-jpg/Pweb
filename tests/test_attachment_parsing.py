from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import xlwt
from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from tender_poc.attachments import AttachmentParser, _extract_text
from tender_poc.storage import TenderStore


class AttachmentParsingTests(unittest.TestCase):
    def test_extracts_pdf_docx_xlsx_and_xls_text(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            pdf_path = root / "sample.pdf"
            docx_path = root / "sample.docx"
            xlsx_path = root / "sample.xlsx"
            xls_path = root / "sample.xls"

            _write_pdf(pdf_path, "spring procurement")

            doc = Document()
            doc.add_paragraph("docx spring procurement")
            doc.save(docx_path)

            workbook = Workbook()
            sheet = workbook.active
            sheet["A1"] = "xlsx spring procurement"
            workbook.save(xlsx_path)

            xls_book = xlwt.Workbook()
            xls_sheet = xls_book.add_sheet("Sheet1")
            xls_sheet.write(0, 0, "xls spring procurement")
            xls_book.save(str(xls_path))

            self.assertIn("spring procurement", _extract_text(pdf_path))
            self.assertIn("docx spring procurement", _extract_text(docx_path))
            self.assertIn("xlsx spring procurement", _extract_text(xlsx_path))
            self.assertIn("xls spring procurement", _extract_text(xls_path))

    def test_doc_without_antiword_records_missing_tool(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            doc_path = root / "sample.doc"
            doc_path.write_bytes(b"not a real doc")
            store = TenderStore(db_path=root / "test.sqlite")
            try:
                parser = AttachmentParser(store)
                with (
                    patch.object(parser, "_download", return_value=doc_path),
                    patch("tender_poc.attachments.shutil.which", return_value=None),
                ):
                    status = parser.parse_one(
                        notice_id="notice-1",
                        attachment_name="sample.doc",
                        attachment_url="https://example.test/sample.doc",
                    )

                self.assertEqual(status, "missing_tool")
                document = store.get_attachment_document("notice-1", "https://example.test/sample.doc")
                self.assertIsNotNone(document)
                self.assertEqual(document["status"], "missing_tool")
            finally:
                store.close()


def _write_pdf(path: Path, text: str) -> None:
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    stream = DecodedStreamObject()
    stream.set_data(f"BT /F1 18 Tf 72 720 Td ({text}) Tj ET".encode("utf-8"))
    page[NameObject("/Contents")] = writer._add_object(stream)
    font = DictionaryObject(
        {
            NameObject("/F1"): DictionaryObject(
                {
                    NameObject("/Type"): NameObject("/Font"),
                    NameObject("/Subtype"): NameObject("/Type1"),
                    NameObject("/BaseFont"): NameObject("/Helvetica"),
                }
            )
        }
    )
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): font})
    with path.open("wb") as fp:
        writer.write(fp)


if __name__ == "__main__":
    unittest.main()
